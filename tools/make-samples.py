"""デモ用のサンプル資料（.pptx）を samples/ に作る。
使い方: python3 tools/make-samples.py
"""
from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.oxml.ns import qn
from pptx.util import Emu, Pt

OUT = Path(__file__).resolve().parent.parent / "samples"
INK = RGBColor(0x1D, 0x1D, 0x1F)
MUTED = RGBColor(0x6E, 0x6E, 0x73)
ACCENT = RGBColor(0x00, 0x71, 0xE3)
ORANGE = RGBColor(0xFF, 0x95, 0x00)


def font(run, size, color, bold=False):
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = color
    run.font.name = "Helvetica Neue"
    rPr = run._r.get_or_add_rPr()
    ea = rPr.makeelement(qn("a:ea"), {})
    rPr.find(qn("a:latin")).addnext(ea)
    ea.set("typeface", "Hiragino Sans")


def text(slide, x, y, w, h, lines, size, color, bold=False):
    tb = slide.shapes.add_textbox(x, y, w, h)
    tf = tb.text_frame
    tf.word_wrap = True
    for i, line in enumerate(lines):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.space_after = Pt(size * 0.6)
        font(p.add_run(), size, color, bold)
        p.runs[0].text = line


def bar(slide, x, y, w, h, color):
    s = slide.shapes.add_shape(1, x, y, w, h)
    s.fill.solid()
    s.fill.fore_color.rgb = color
    s.line.fill.background()


def deck(path, width_in, height_in, slides, accent):
    prs = Presentation()
    prs.slide_width = Emu(int(width_in * 914400))
    prs.slide_height = Emu(int(height_in * 914400))
    W, H = prs.slide_width, prs.slide_height
    m = int(W * 0.08)
    for n, (title, body) in enumerate(slides, 1):
        s = prs.slides.add_slide(prs.slide_layouts[6])
        if n == 1:
            text(s, m, int(H * 0.36), W - 2 * m, int(H * 0.2), [title], 40, INK, True)
            bar(s, m, int(H * 0.56), int(W * 0.08), Emu(60000), accent)
            text(s, m, int(H * 0.62), W - 2 * m, int(H * 0.2), body, 20, MUTED)
        else:
            text(s, m, int(H * 0.09), W - 2 * m, int(H * 0.14), [title], 32, INK, True)
            bar(s, m, int(H * 0.24), int(W * 0.08), Emu(50000), accent)
            text(s, m, int(H * 0.31), W - 2 * m, int(H * 0.6), body, 20, INK)
        text(s, W - m - Emu(914400), H - int(H * 0.1), Emu(914400), int(H * 0.06), [f"{n}"], 12, MUTED)
    prs.save(path)


OUT.mkdir(exist_ok=True)

deck(OUT / "usage.pptx", 13.333, 7.5, [
    ("slide-viewer", ["Web ページに埋め込むスライドビューア"]),
    ("埋め込み方", ["script タグを1行読み込む", "表示したい場所に div を1つ置く", "data-deck に資料名を書く"]),
    ("ページのめくり方", ["← → キー", "スワイプ", "前へ・次へボタン、サムネイル"]),
    ("ダウンロードの設定", ["資料ごとに deck.json で許可・不許可を決める", "ページ側でボタンを隠すこともできる", "不許可の資料は PDF をサーバーに置かない"]),
    ("資料の登録", ["tools/add-deck.sh 資料.pptx 資料名", "PDF と PowerPoint のどちらでもよい", "--download を付けると PDF も置く"]),
    ("おわり", ["このサンプルはダウンロードできる設定"]),
], ACCENT)

deck(OUT / "locked.pptx", 10, 7.5, [
    ("4:3 のサンプル", ["ダウンロードできない設定の資料"]),
    ("縦横比", ["資料ごとの縦横比に合わせて表示する", "16:9 でも 4:3 でも崩れない"]),
    ("守っているもの", ["元の PDF はサーバーに置いていない", "ページ画像は短時間だけ有効なトークンで返す", "URL を別のブラウザに貼っても開けない"]),
    ("防げないもの", ["スクリーンショット", "表示中の画像を1枚ずつ保存すること"]),
    ("おわり", ["このサンプルはダウンロードできない設定"]),
], ORANGE)

print("samples/usage.pptx samples/locked.pptx")
